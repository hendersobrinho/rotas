"""Cadastro de clientes a partir de uma planilha, e o modelo para preencher.

A planilha tem uma linha por cliente. Os dados do cliente ocupam as primeiras
colunas; cada endereço ocupa um bloco de colunas próprio — "Endereço 1 - ...",
"Endereço 2 - ..." — porque um cliente pode ter quantos endereços precisar e
numa linha só não caberia de outro jeito.

O modelo sai com três blocos, que dá conta da maioria. Quem tiver um cliente
com mais é só copiar o último bloco e continuar a numeração: a leitura
descobre quantos blocos existem pelo cabeçalho, não por um número fixo.

O cabeçalho é reconhecido sem frescura de acento, caixa ou pontuação, então
uma planilha que o escritório já mantém costuma entrar sem ser reescrita.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app import planilha
from app.models import Cliente, TipoCliente, TipoEndereco
from app.repository import clientes as repo_clientes
from app.schemas import DadosCliente, DadosEndereco

# (campo do schema, título da coluna) — a ordem é a ordem da planilha.
CAMPOS_CLIENTE = (
    ("tipo", "Tipo (PF ou PJ)"),
    ("nome", "Nome / razão social"),
    ("apelido", "Apelido / nome social"),
    ("telefone", "Telefone"),
    ("observacao", "Observação"),
)
CAMPOS_ENDERECO = (
    ("tipo", "Tipo"),
    ("rotulo", "Nome do endereço"),
    ("logradouro", "Logradouro"),
    ("numero", "Número"),
    ("complemento", "Complemento"),
    ("bairro", "Bairro"),
    ("cidade", "Cidade"),
    ("cep", "CEP"),
    ("observacao", "Observação"),
)
BLOCOS_NO_MODELO = 3

# O banco corta o que passar disso; melhor avisar na conferência do que deixar
# o erro estourar no meio da gravação. (Ver app/models.py.)
LIMITES_CLIENTE = {"nome": 200, "apelido": 120, "telefone": 40}
LIMITES_ENDERECO = {
    "rotulo": 60,
    "logradouro": 200,
    "numero": 20,
    "complemento": 100,
    "bairro": 100,
    "cidade": 100,
    "cep": 15,
}


def _chave(texto: str) -> str:
    """Texto comparável: sem acento, sem caixa, sem pontuação, sem espaço extra."""
    decomposto = unicodedata.normalize("NFKD", texto or "")
    sem_acento = "".join(c for c in decomposto if not unicodedata.combining(c))
    return " ".join(re.sub(r"[^A-Za-z0-9]+", " ", sem_acento).lower().split())


_SINONIMOS_CLIENTE = {
    "tipo": {"tipo", "tipo pf ou pj", "tipo de cliente", "pf ou pj", "pessoa"},
    "nome": {"nome razao social", "nome", "razao social", "nome completo", "cliente"},
    "apelido": {"apelido nome social", "apelido", "nome social", "nome fantasia"},
    "telefone": {"telefone", "telefones", "fone", "celular", "contato"},
    "observacao": {"observacao", "observacoes", "obs"},
}
_SINONIMOS_ENDERECO = {
    "tipo": {"tipo", "tipo de endereco"},
    "rotulo": {"nome do endereco", "nome", "rotulo", "apelido", "identificacao"},
    "logradouro": {"logradouro", "rua", "endereco", "avenida"},
    "numero": {"numero", "num", "n", "nro"},
    "complemento": {"complemento", "compl"},
    "bairro": {"bairro"},
    "cidade": {"cidade", "municipio"},
    "cep": {"cep"},
    "observacao": {"observacao do endereco", "observacao", "observacoes", "obs"},
}

_TIPOS_CLIENTE = {
    "pf": TipoCliente.PF, "f": TipoCliente.PF, "fisica": TipoCliente.PF,
    "pessoa fisica": TipoCliente.PF, "pessoa f": TipoCliente.PF,
    "pj": TipoCliente.PJ, "j": TipoCliente.PJ, "juridica": TipoCliente.PJ,
    "pessoa juridica": TipoCliente.PJ, "pessoa j": TipoCliente.PJ,
    "empresa": TipoCliente.PJ,
}
_TIPOS_ENDERECO = {
    "residencial": TipoEndereco.RESIDENCIAL, "residencia": TipoEndereco.RESIDENCIAL,
    "res": TipoEndereco.RESIDENCIAL, "r": TipoEndereco.RESIDENCIAL,
    "casa": TipoEndereco.RESIDENCIAL,
    "comercial": TipoEndereco.COMERCIAL, "com": TipoEndereco.COMERCIAL,
    "c": TipoEndereco.COMERCIAL, "empresa": TipoEndereco.COMERCIAL,
    "trabalho": TipoEndereco.COMERCIAL, "escritorio": TipoEndereco.COMERCIAL,
}

_PADRAO_BLOCO = re.compile(r"^enderecos? (\d+)\s*(.*)$")


def cabecalho(blocos: int = BLOCOS_NO_MODELO) -> list[str]:
    """Os títulos das colunas, na ordem: cliente e depois um bloco por endereço."""
    titulos = [titulo for _campo, titulo in CAMPOS_CLIENTE]
    for numero in range(1, blocos + 1):
        titulos += [f"Endereço {numero} - {titulo}" for _campo, titulo in CAMPOS_ENDERECO]
    return titulos


# --------------------------------------------------------------- o que se lê
@dataclass
class LinhaImportacao:
    """Uma linha da planilha já entendida — ou o motivo de não ter dado."""

    numero: int                       # a linha como ela aparece na planilha
    dados: DadosCliente | None = None
    problema: str | None = None
    existente_id: int | None = None   # o cliente já cadastrado com esse nome
    tipo_entendido: bool = True       # False quando a coluna PF/PJ não serviu

    @property
    def valida(self) -> bool:
        return self.dados is not None and self.problema is None

    @property
    def nome(self) -> str:
        return self.dados.nome if self.dados else "—"

    @property
    def rotulo_tipo(self) -> str:
        """PF/PJ para mostrar — travessão quando a coluna não deu para ler.

        Sem isto a prévia mostraria “Pessoa Física” na mesma linha em que diz
        que o tipo não foi entendido, que é o oposto de ajudar.
        """
        if self.dados is None or not self.tipo_entendido:
            return "—"
        return self.dados.tipo.value

    @property
    def quantos_enderecos(self) -> int:
        return len(self.dados.enderecos) if self.dados else 0

    def situacao(self, atualizar: bool) -> str:
        """O que vai acontecer com esta linha, em uma palavra."""
        if self.problema:
            return self.problema
        if self.existente_id is None:
            return "Novo cadastro"
        return "Atualiza o cadastro" if atualizar else "Já cadastrado — fica de fora"

    def sera_gravada(self, atualizar: bool) -> bool:
        return self.valida and (self.existente_id is None or atualizar)


@dataclass
class Resultado:
    criados: int = 0
    atualizados: int = 0
    ignorados: int = 0

    def resumo(self) -> str:
        partes = []
        if self.criados:
            partes.append(f"{self.criados} cliente(s) cadastrado(s)")
        if self.atualizados:
            partes.append(f"{self.atualizados} atualizado(s)")
        if self.ignorados:
            partes.append(f"{self.ignorados} deixado(s) de fora")
        return ", ".join(partes) or "Nada foi gravado."


# ------------------------------------------------------------------- leitura
def _mapear_colunas(titulos: list[str]) -> tuple[dict[str, int], dict[int, dict[str, int]]]:
    """Descobre em que coluna está cada campo, pelo cabeçalho."""
    do_cliente: dict[str, int] = {}
    dos_blocos: dict[int, dict[str, int]] = {}

    for indice, titulo in enumerate(titulos):
        chave = _chave(titulo)
        if not chave:
            continue

        bloco = _PADRAO_BLOCO.match(chave)
        if bloco is not None:
            numero = int(bloco.group(1))
            resto = bloco.group(2)
            for campo, sinonimos in _SINONIMOS_ENDERECO.items():
                if resto in sinonimos:
                    dos_blocos.setdefault(numero, {}).setdefault(campo, indice)
                    break
            continue

        for campo, sinonimos in _SINONIMOS_CLIENTE.items():
            if chave in sinonimos:
                do_cliente.setdefault(campo, indice)
                break

    return do_cliente, dos_blocos


def _achar_cabecalho(linhas: list[list[str]]) -> int:
    """A linha do cabeçalho — costuma ser a primeira, mas nem sempre é.

    Vale a primeira que tenha a coluna do nome: acima dela pode haver título,
    data de extração e outras sobras que quem exporta de outro sistema deixa.
    """
    for indice, linha in enumerate(linhas[:20]):
        do_cliente, _ = _mapear_colunas(linha)
        if "nome" in do_cliente:
            return indice
    raise ValueError(
        "Não achei o cabeçalho na planilha. Ela precisa de uma linha com os "
        "títulos das colunas — a coluna “Nome / razão social” é obrigatória.\n\n"
        "Use o botão “Baixar modelo” para partir de uma planilha já pronta."
    )


def _valor(linha: list[str], indice: int | None) -> str:
    if indice is None or indice >= len(linha):
        return ""
    return (linha[indice] or "").strip()


def _endereco_da_linha(
    linha: list[str], colunas: dict[str, int], numero: int, problemas: list[str]
) -> DadosEndereco | None:
    """Monta um endereço de um bloco de colunas; None se o bloco está vazio."""
    bruto = {campo: _valor(linha, colunas.get(campo)) for campo in _SINONIMOS_ENDERECO}
    if not any(bruto.values()):
        return None

    texto_tipo = _chave(bruto["tipo"])
    if not texto_tipo:
        # O escritório atende quase só no comercial; exigir o tipo em cada
        # bloco preenchido só faria a planilha ser devolvida por detalhe.
        tipo = TipoEndereco.COMERCIAL
    elif texto_tipo in _TIPOS_ENDERECO:
        tipo = _TIPOS_ENDERECO[texto_tipo]
    else:
        problemas.append(
            f"endereço {numero}: “{bruto['tipo']}” não é Residencial nem Comercial"
        )
        return None

    for campo, limite in LIMITES_ENDERECO.items():
        if len(bruto[campo]) > limite:
            problemas.append(
                f"endereço {numero}: {campo} passa de {limite} caracteres"
            )

    return DadosEndereco(
        tipo=tipo,
        rotulo=bruto["rotulo"] or None,
        logradouro=bruto["logradouro"] or None,
        numero=bruto["numero"] or None,
        complemento=bruto["complemento"] or None,
        bairro=bruto["bairro"] or None,
        cidade=bruto["cidade"] or None,
        cep=bruto["cep"] or None,
        observacao=bruto["observacao"] or None,
    )


def _cliente_da_linha(
    linha: list[str],
    do_cliente: dict[str, int],
    dos_blocos: dict[int, dict[str, int]],
) -> tuple[DadosCliente | None, bool, str | None]:
    problemas: list[str] = []

    nome = _valor(linha, do_cliente.get("nome"))
    if not nome:
        return None, False, "sem nome"

    texto_tipo = _chave(_valor(linha, do_cliente.get("tipo")))
    tipo_entendido = texto_tipo in _TIPOS_CLIENTE
    if tipo_entendido:
        tipo = _TIPOS_CLIENTE[texto_tipo]
    else:
        # PF é só para o objeto ficar montável; a linha não será gravada.
        tipo = TipoCliente.PF
        problemas.append(
            "falta dizer se é PF ou PJ" if not texto_tipo
            else f"tipo “{_valor(linha, do_cliente.get('tipo'))}” não é PF nem PJ"
        )

    for campo, limite in LIMITES_CLIENTE.items():
        if len(_valor(linha, do_cliente.get(campo))) > limite:
            problemas.append(f"{campo} passa de {limite} caracteres")

    enderecos: list[DadosEndereco] = []
    for numero in sorted(dos_blocos):
        endereco = _endereco_da_linha(linha, dos_blocos[numero], numero, problemas)
        if endereco is not None:
            enderecos.append(endereco)

    dados = DadosCliente(
        tipo=tipo,
        nome=nome,
        apelido=_valor(linha, do_cliente.get("apelido")) or None,
        telefone=_valor(linha, do_cliente.get("telefone")) or None,
        observacao=_valor(linha, do_cliente.get("observacao")) or None,
        enderecos=enderecos,
    ).normalizado()

    if not problemas:
        # As mesmas regras da tela — dois endereços com o mesmo nome, por
        # exemplo, o repositório recusa na hora de gravar.
        try:
            repo_clientes.validar(dados)
        except ValueError as erro:
            problemas.append(str(erro))

    return dados, tipo_entendido, "; ".join(problemas) if problemas else None


def ler(caminho: str | Path) -> list[LinhaImportacao]:
    """Lê a planilha e devolve uma entrada por linha preenchida."""
    conteudo = planilha.ler(caminho)
    if not conteudo:
        raise ValueError("A planilha está vazia.")

    inicio = _achar_cabecalho(conteudo)
    do_cliente, dos_blocos = _mapear_colunas(conteudo[inicio])

    lidas: list[LinhaImportacao] = []
    for deslocamento, linha in enumerate(conteudo[inicio + 1:], start=inicio + 2):
        if not any((celula or "").strip() for celula in linha):
            continue  # linha em branco no meio da planilha não é erro
        dados, tipo_ok, problema = _cliente_da_linha(linha, do_cliente, dos_blocos)
        lidas.append(LinhaImportacao(
            numero=deslocamento, dados=dados, problema=problema,
            tipo_entendido=tipo_ok,
        ))

    if not lidas:
        raise ValueError("A planilha tem o cabeçalho, mas nenhuma linha preenchida.")
    return lidas


def conferir(sessao: Session, linhas: list[LinhaImportacao]) -> list[LinhaImportacao]:
    """Marca quem já está cadastrado e quem se repete dentro da própria planilha."""
    nomes = {
        _chave(linha.dados.nome): linha for linha in linhas if linha.dados is not None
    }
    if nomes:
        consulta = select(Cliente).where(
            func.lower(Cliente.nome).in_([
                linha.dados.nome.lower() for linha in linhas if linha.dados
            ])
        )
        for cliente in sessao.scalars(consulta):
            alvo = nomes.get(_chave(cliente.nome))
            if alvo is not None:
                alvo.existente_id = cliente.id

    vistos: dict[str, int] = {}
    for linha in linhas:
        if linha.dados is None:
            continue
        chave = _chave(linha.dados.nome)
        primeira = vistos.get(chave)
        if primeira is None:
            vistos[chave] = linha.numero
        elif not linha.problema:
            linha.problema = f"nome repetido (já aparece na linha {primeira})"
    return linhas


def analisar(sessao: Session, caminho: str | Path) -> list[LinhaImportacao]:
    """Lê a planilha e já confere contra o que está no banco."""
    return conferir(sessao, ler(caminho))


def importar(
    sessao: Session, linhas: list[LinhaImportacao], atualizar_existentes: bool = False
) -> Resultado:
    """Grava as linhas boas. As demais só são contadas."""
    resultado = Resultado()
    for linha in linhas:
        if not linha.sera_gravada(atualizar_existentes):
            resultado.ignorados += 1
            continue
        if linha.existente_id is None:
            repo_clientes.criar_cliente(sessao, linha.dados)
            resultado.criados += 1
        else:
            repo_clientes.atualizar_cliente(sessao, linha.existente_id, linha.dados)
            resultado.atualizados += 1
    return resultado


# -------------------------------------------------------- modelo em branco
LARGURAS = {
    "Tipo (PF ou PJ)": 14,
    "Nome / razão social": 38,
    "Apelido / nome social": 24,
    "Telefone": 18,
    "Observação": 34,
    "Tipo": 14,
    "Nome do endereço": 20,
    "Logradouro": 30,
    "Número": 10,
    "Complemento": 18,
    "Bairro": 20,
    "Cidade": 20,
    "CEP": 12,
}
LINHAS_PREPARADAS = 300  # até onde o modelo já vai formatado para preencher

# As cores da marca, repetidas aqui de propósito: este módulo não é de tela e
# não deve arrastar o PySide6 junto só para pintar um cabeçalho de planilha.
# Se a marca mudar, o original está em app/ui/marca.py.
AZUL_CABECALHO = "203461"
CIANO_CABECALHO = "4BBDCD"

INSTRUCOES = (
    ("Tipo (PF ou PJ)", "Sim",
     "PF para pessoa física, PJ para pessoa jurídica. Também aceita "
     "“Pessoa Física” e “Pessoa Jurídica” por extenso."),
    ("Nome / razão social", "Sim",
     "Razão social da empresa ou o nome completo da pessoa. É por ele que o "
     "sistema reconhece um cliente que já existe."),
    ("Apelido / nome social", "Não",
     "Como o escritório chama o cliente no dia a dia. É o que aparece nas "
     "listas e na agenda quando está preenchido."),
    ("Telefone", "Não", "Um telefone de contato, do jeito que for usado."),
    ("Observação", "Não",
     "Combinados, recados, horários — texto livre."),
    ("Endereço N - Tipo", "Não",
     "Residencial ou Comercial. Em branco, entra como Comercial."),
    ("Endereço N - Nome do endereço", "Não",
     "Como este endereço é chamado: Matriz, Filial Barreiro, Casa da Ana. "
     "Obrigatório quando o cliente tem mais de um endereço do mesmo tipo, "
     "senão não há como diferenciar um do outro na hora de marcar o serviço."),
    ("Endereço N - Logradouro até CEP", "Não",
     "Os dados do endereço. Um bloco totalmente em branco é ignorado."),
    ("Endereço N - Observação", "Não",
     "Referência, portaria, horário de entrega — o que for do endereço."),
)


def gerar_modelo(caminho: str | Path, blocos: int = BLOCOS_NO_MODELO) -> Path:
    """Escreve a planilha em branco, pronta para preencher e reimportar."""
    caminho = Path(caminho)
    titulos = cabecalho(blocos)
    if caminho.suffix.lower() == ".csv" or not planilha.TEM_XLSX:
        planilha.escrever_csv(caminho, [titulos])
        return caminho
    _modelo_xlsx(caminho, titulos)
    return caminho


def _modelo_xlsx(caminho: Path, titulos: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    livro = Workbook()
    aba = livro.active
    aba.title = "Clientes"

    fundo_cliente = PatternFill("solid", fgColor=AZUL_CABECALHO)
    fundo_endereco = PatternFill("solid", fgColor=CIANO_CABECALHO)
    borda = Border(bottom=Side(style="thin", color="FFFFFF"))

    for coluna, titulo in enumerate(titulos, start=1):
        celula = aba.cell(row=1, column=coluna, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF", size=11)
        celula.fill = fundo_endereco if titulo.startswith("Endereço") else fundo_cliente
        celula.alignment = Alignment(vertical="center", wrap_text=True)
        celula.border = borda
        largura = LARGURAS.get(titulo.split(" - ")[-1], 26)
        aba.column_dimensions[get_column_letter(coluna)].width = largura

    aba.row_dimensions[1].height = 32
    aba.freeze_panes = "B2"  # o cabeçalho e a coluna do tipo ficam à vista

    # Tudo como texto: senão o Excel transforma um CEP "30140002" em número e
    # engole o zero à esquerda de quem tem.
    for linha in range(2, LINHAS_PREPARADAS + 2):
        for coluna in range(1, len(titulos) + 1):
            aba.cell(row=linha, column=coluna).number_format = "@"

    intervalo_ate = LINHAS_PREPARADAS + 1
    for coluna, titulo in enumerate(titulos, start=1):
        if not titulo.endswith("Tipo") and titulo != "Tipo (PF ou PJ)":
            continue
        opcoes = '"PF,PJ"' if coluna == 1 else '"Comercial,Residencial"'
        validacao = DataValidation(type="list", formula1=opcoes, allow_blank=True)
        validacao.error = "Escolha uma das opções da lista."
        validacao.prompt = "PF ou PJ" if coluna == 1 else "Comercial ou Residencial"
        aba.add_data_validation(validacao)
        letra = get_column_letter(coluna)
        validacao.add(f"{letra}2:{letra}{intervalo_ate}")

    ajuda = livro.create_sheet("Como preencher")
    ajuda.column_dimensions["A"].width = 34
    ajuda.column_dimensions["B"].width = 12
    ajuda.column_dimensions["C"].width = 78

    ajuda["A1"] = "Como preencher a planilha de clientes"
    ajuda["A1"].font = Font(bold=True, size=14, color=AZUL_CABECALHO)
    ajuda["A2"] = (
        "Uma linha por cliente, na aba “Clientes”. Não mude os títulos das "
        "colunas: é por eles que o sistema entende a planilha."
    )
    ajuda["A3"] = (
        "Cada endereço tem o seu bloco de colunas. Precisa de um quarto "
        "endereço? Copie as colunas do “Endereço 3”, cole no fim e troque o "
        "número para 4 — a importação conta os blocos pelo cabeçalho."
    )
    ajuda["A4"] = (
        "Cliente cujo nome já está cadastrado fica de fora, a não ser que a "
        "opção de atualizar esteja marcada na hora de importar."
    )
    for linha in (2, 3, 4):
        ajuda[f"A{linha}"].alignment = Alignment(wrap_text=True, vertical="top")
        ajuda.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=3)
        ajuda.row_dimensions[linha].height = 30

    for coluna, titulo in enumerate(("Coluna", "Obrigatório", "O que vai nela"), start=1):
        celula = ajuda.cell(row=6, column=coluna, value=titulo)
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = fundo_cliente

    for indice, (nome, obrigatorio, explicacao) in enumerate(INSTRUCOES, start=7):
        ajuda.cell(row=indice, column=1, value=nome).alignment = Alignment(
            vertical="top"
        )
        ajuda.cell(row=indice, column=2, value=obrigatorio).alignment = Alignment(
            horizontal="center", vertical="top"
        )
        celula = ajuda.cell(row=indice, column=3, value=explicacao)
        celula.alignment = Alignment(wrap_text=True, vertical="top")
        ajuda.row_dimensions[indice].height = 30

    livro.save(caminho)
