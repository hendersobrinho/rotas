"""Importação de clientes por planilha: modelo, leitura, conferência e gravação."""

from __future__ import annotations

import sys
from pathlib import Path

from comum import RAIZ, preparar_ambiente, semear

preparar_ambiente()

from PySide6.QtWidgets import QApplication  # noqa: E402

app = QApplication([])

import openpyxl  # noqa: E402

from app import importacao_clientes as importacao  # noqa: E402
from app import planilha  # noqa: E402
from app.db import session_scope  # noqa: E402
from app.models import TipoCliente, TipoEndereco  # noqa: E402
from app.repository import clientes as repo_clientes  # noqa: E402
from app.ui.estilo import aplicar_tema  # noqa: E402
from app.ui.importar_clientes import ImportarClientesDialog  # noqa: E402

SAIDA = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "testes" / ".saida"
SAIDA.mkdir(parents=True, exist_ok=True)

aplicar_tema(app)
semear(com_eventos=False)


def escrever(nome: str, linhas: list[list], blocos: int = 3) -> Path:
    """Gera o modelo e preenche com as linhas pedidas, como o usuário faria."""
    caminho = SAIDA / nome
    importacao.gerar_modelo(caminho, blocos=blocos)
    if caminho.suffix == ".csv":
        conteudo = planilha.ler(caminho) + [[str(c) for c in l] for l in linhas]
        planilha.escrever_csv(caminho, conteudo)
        return caminho
    livro = openpyxl.load_workbook(caminho)
    aba = livro["Clientes"]
    # O modelo já vem com as linhas formatadas, então escreve-se por posição
    # em vez de append(), que cairia depois delas.
    for deslocamento, valores in enumerate(linhas, start=2):
        for coluna, valor in enumerate(valores, start=1):
            aba.cell(row=deslocamento, column=coluna, value=valor)
    livro.save(caminho)
    return caminho


# ---- o modelo ------------------------------------------------------------
modelo = SAIDA / "modelo.xlsx"
importacao.gerar_modelo(modelo)
livro = openpyxl.load_workbook(modelo)
assert livro.sheetnames == ["Clientes", "Como preencher"], livro.sheetnames
titulos = [c.value for c in livro["Clientes"][1]]
assert titulos == importacao.cabecalho(3), titulos
assert titulos[:5] == [
    "Tipo (PF ou PJ)", "Nome / razão social", "Apelido / nome social",
    "Telefone", "Observação",
], titulos[:5]
# cada endereço em colunas próprias, os três blocos completos
for numero in (1, 2, 3):
    for _campo, nome_campo in importacao.CAMPOS_ENDERECO:
        assert f"Endereço {numero} - {nome_campo}" in titulos, (numero, nome_campo)
assert len(titulos) == 5 + 3 * len(importacao.CAMPOS_ENDERECO), len(titulos)
# lista suspensa no tipo, para ninguém digitar "PF." e o sistema recusar
validacoes = livro["Clientes"].data_validations.dataValidation
assert len(validacoes) == 4, len(validacoes)
print(f"ok modelo com {len(titulos)} colunas e 3 blocos de endereço")

importacao.gerar_modelo(SAIDA / "modelo.csv")
assert planilha.ler(SAIDA / "modelo.csv")[0] == importacao.cabecalho(3)
print("ok modelo em csv")

# ---- leitura de uma planilha preenchida ----------------------------------
LINHAS = [
    # PJ com dois endereços, cada um no seu bloco
    ["PJ", "Contabilidade Nova Era LTDA", "Nova Era", "(31) 3010-2020", "Só de manhã",
     "Comercial", "Matriz", "Rua Bahia", "1500", "Sala 704", "Centro",
     "Belo Horizonte", "30160-011", "Portaria fecha 18h",
     "Residencial", "Casa do Dr. Paulo", "Rua Pium-í", "300", "Apto 502",
     "Cruzeiro", "Belo Horizonte", "30310-080", ""],
    # PF sem endereço nenhum, e o tipo escrito por extenso
    ["Pessoa Física", "Marina Alves Rocha", "Marina", "(31) 99777-1234", "", ],
    # três endereços: o terceiro bloco também é lido
    ["pj", "Mercado Dois Irmãos ME", "Dois Irmãos", "", "",
     "comercial", "Loja Centro", "Rua Curitiba", "90", "", "Centro", "BH", "", "",
     "COMERCIAL", "Loja Barreiro", "Av. Olinto", "12", "", "Barreiro", "BH", "", "",
     "residencial", "Casa do Zé", "Rua das Flores", "7", "", "Eldorado",
     "Contagem", "32310-100", "Cachorro no portão"],
    # bloco sem tipo entra como comercial; CEP digitado como número
    ["PJ", "Transportadora Vale Verde", "Vale Verde", "", "",
     "", "Galpão", "Rodovia BR-040", "km 8", "", "Distrito", "Betim", 32600010, ""],
]
caminho = escrever("clientes.xlsx", LINHAS)
lidas = importacao.ler(caminho)
assert len(lidas) == 4, [l.nome for l in lidas]
assert all(l.valida for l in lidas), [(l.nome, l.problema) for l in lidas]

por_nome = {l.nome: l for l in lidas}
nova_era = por_nome["Contabilidade Nova Era LTDA"].dados
assert nova_era.tipo is TipoCliente.PJ
assert nova_era.apelido == "Nova Era" and nova_era.telefone == "(31) 3010-2020"
assert len(nova_era.enderecos) == 2, len(nova_era.enderecos)
assert nova_era.enderecos[0].tipo is TipoEndereco.COMERCIAL
assert nova_era.enderecos[0].complemento == "Sala 704"
assert nova_era.enderecos[1].tipo is TipoEndereco.RESIDENCIAL
assert nova_era.enderecos[1].rotulo == "Casa do Dr. Paulo"
assert por_nome["Marina Alves Rocha"].dados.tipo is TipoCliente.PF
assert por_nome["Marina Alves Rocha"].quantos_enderecos == 0
assert por_nome["Mercado Dois Irmãos ME"].quantos_enderecos == 3
vale = por_nome["Transportadora Vale Verde"].dados
assert vale.enderecos[0].tipo is TipoEndereco.COMERCIAL, "bloco sem tipo vira comercial"
assert vale.enderecos[0].cep == "32600010", vale.enderecos[0].cep
print("ok leu 4 clientes, com 0, 1, 2 e 3 endereços")

# ---- gravação ------------------------------------------------------------
with session_scope() as sessao:
    conferidas = importacao.conferir(sessao, lidas)
    assert all(l.existente_id is None for l in conferidas)
    resultado = importacao.importar(sessao, conferidas)
assert (resultado.criados, resultado.atualizados, resultado.ignorados) == (4, 0, 0), resultado

with session_scope() as sessao:
    gravado = repo_clientes.listar_clientes(sessao, "Nova Era")[0]
    assert gravado.nome == "Contabilidade Nova Era LTDA"
    assert [e.etiqueta for e in gravado.enderecos] == ["Matriz", "Casa do Dr. Paulo"]
    assert gravado.enderecos[0].cidade == "Belo Horizonte"
    irmaos = repo_clientes.listar_clientes(sessao, "Dois Irmãos")[0]
    assert len(irmaos.enderecos) == 3, len(irmaos.enderecos)
print("ok gravou os 4 com os endereços em colunas separadas")

# ---- quem já existe ------------------------------------------------------
REPETIDA = [
    ["PJ", "Contabilidade Nova Era LTDA", "Nova Era do Sul", "(31) 3010-9999", "",
     "Comercial", "Matriz nova", "Av. Afonso Pena", "1", "", "Centro", "BH", "", ""],
]
caminho = escrever("repetida.xlsx", REPETIDA)
with session_scope() as sessao:
    linhas = importacao.analisar(sessao, caminho)
    assert linhas[0].existente_id is not None, "não reconheceu o cliente já cadastrado"
    assert not linhas[0].sera_gravada(atualizar=False)
    assert "Já cadastrado" in linhas[0].situacao(atualizar=False)
    resultado = importacao.importar(sessao, linhas, atualizar_existentes=False)
assert (resultado.criados, resultado.ignorados) == (0, 1), resultado

with session_scope() as sessao:
    linhas = importacao.analisar(sessao, caminho)
    resultado = importacao.importar(sessao, linhas, atualizar_existentes=True)
assert (resultado.criados, resultado.atualizados) == (0, 1), resultado
with session_scope() as sessao:
    atual = repo_clientes.listar_clientes(sessao, "Nova Era")[0]
    assert atual.apelido == "Nova Era do Sul", atual.apelido
    assert [e.etiqueta for e in atual.enderecos] == ["Matriz nova"], atual.enderecos
print("ok cliente repetido: fica de fora, ou é atualizado quando se pede")

# ---- linhas com problema -------------------------------------------------
TORTAS = [
    ["XPTO", "Cliente com tipo errado"],
    ["PF", ""],                                   # sem nome
    ["PJ", "Padaria Duas Portas", "", "", "",
     "Comercial", "Loja", "Rua A", "1", "", "", "", "", "",
     "Comercial", "loja", "Rua B", "2", "", "", "", "", ""],   # rótulos iguais
    ["PJ", "x" * 250],                            # passa do limite do banco
    ["PJ", "Cliente Bom LTDA", "", "", "",
     "Comercial", "Sede", "Rua Boa", "10", "", "Centro", "BH", "", ""],
    ["PJ", "Cliente Bom LTDA"],                   # repetido dentro da planilha
]
caminho = escrever("tortas.xlsx", TORTAS)
with session_scope() as sessao:
    linhas = importacao.analisar(sessao, caminho)
problemas = {l.numero: l.problema for l in linhas}
assert "não é PF nem PJ" in problemas[2], problemas
assert problemas[3] == "sem nome", problemas
assert "dois endereços chamados" in problemas[4], problemas
assert "200 caracteres" in problemas[5], problemas
assert problemas[6] is None, problemas
assert "repetido" in problemas[7], problemas
boas = [l for l in linhas if l.sera_gravada(atualizar=False)]
assert len(boas) == 1 and boas[0].nome == "Cliente Bom LTDA"
with session_scope() as sessao:
    resultado = importacao.importar(sessao, linhas)
assert (resultado.criados, resultado.ignorados) == (1, 5), resultado
print("ok 5 linhas tortas ficaram de fora e a boa entrou")

# ---- cabeçalho de outro jeito, e csv -------------------------------------
OUTRO = [
    ["tipo", "RAZÃO SOCIAL", "nome fantasia", "fone",
     "endereco 1 tipo", "Endereco 1 Rua", "endereço 1 - numero", "Endereço 1 Cidade"],
    ["PJ", "Livraria do Beco ME", "Livraria do Beco", "3133334444",
     "Comercial", "Rua dos Caetés", "55", "Belo Horizonte"],
]
caminho = SAIDA / "outro_cabecalho.csv"
planilha.escrever_csv(caminho, OUTRO)
linhas = importacao.ler(caminho)
assert len(linhas) == 1 and linhas[0].valida, [l.problema for l in linhas]
assert linhas[0].dados.apelido == "Livraria do Beco"
assert linhas[0].dados.enderecos[0].logradouro == "Rua dos Caetés"
assert linhas[0].dados.enderecos[0].cidade == "Belo Horizonte"
print("ok cabeçalho sem acento/caixa/pontuação e planilha em csv")

# ---- planilhas que não servem -------------------------------------------
vazia = SAIDA / "vazia.csv"
planilha.escrever_csv(vazia, [["Data da extração"], ["nada aqui"]])
try:
    importacao.ler(vazia)
    raise AssertionError("aceitou planilha sem cabeçalho")
except ValueError as erro:
    assert "cabeçalho" in str(erro), erro

so_cabecalho = SAIDA / "so_cabecalho.xlsx"
importacao.gerar_modelo(so_cabecalho)
try:
    importacao.ler(so_cabecalho)
    raise AssertionError("aceitou planilha sem nenhuma linha")
except ValueError as erro:
    assert "nenhuma linha" in str(erro), erro
print("ok planilha sem cabeçalho e planilha vazia são recusadas")

# ---- a tela --------------------------------------------------------------
dialogo = ImportarClientesDialog()
assert not dialogo.btn_importar.isEnabled(), "importar ligado sem planilha escolhida"
caminho = escrever("para_tela.xlsx", [
    ["PJ", "Ótica Olhar Certo ME", "Olhar Certo", "", "",
     "Comercial", "Loja", "Rua Rio de Janeiro", "400", "", "Centro", "BH", "", ""],
    ["PJ", "Cliente Bom LTDA"],   # já cadastrado pelo bloco anterior
    ["ZZZ", "Tipo inválido"],
])
with session_scope() as sessao:
    dialogo._linhas = importacao.analisar(sessao, caminho)
dialogo._caminho = caminho
dialogo._mostrar_previa()

assert dialogo.tabela.rowCount() == 3, dialogo.tabela.rowCount()
assert dialogo.btn_importar.isEnabled()
situacoes = [dialogo.tabela.item(l, 4).text() for l in range(3)]
assert situacoes[0] == "Novo cadastro", situacoes
assert "Já cadastrado" in situacoes[1], situacoes
assert "não é PF nem PJ" in situacoes[2], situacoes
assert "1 novo(s)" in dialogo.resumo.text(), dialogo.resumo.text()
assert "1 com problema" in dialogo.resumo.text(), dialogo.resumo.text()

dialogo.campo_atualizar.setChecked(True)
assert "Atualiza o cadastro" in dialogo.tabela.item(1, 4).text()
print("ok prévia da tela mostra novo, já cadastrado e problema")

dialogo._linhas = []
dialogo._caminho = None
dialogo._mostrar_previa()
assert not dialogo.btn_importar.isEnabled()

print("TESTE IMPORTACAO OK")
